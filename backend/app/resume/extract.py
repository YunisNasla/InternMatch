"""Resume text from an uploaded PDF or Word file, in memory only. The file is never written to disk."""

import io
import zipfile
from pathlib import Path

import docx
import pdfplumber

MAX_BYTES = 5 * 1024 * 1024
# shortcut: resumes are 1 to 2 pages; these caps stop a huge PDF from stalling the server or the parser call. Raise them if real CVs get cut off.
MAX_PAGES = 6
MAX_CHARS = 30_000
# A Word file is a zip. Its XML is parsed into a full tree (about 50x its size in memory), so it gets a tight cap;
# images and other media only count toward the total.
MAX_UNPACKED = 50 * 1024 * 1024
MAX_XML = 5 * 1024 * 1024


class _TooLarge(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    """Plain text of the resume. Raises ValueError with a message the user can act on."""
    if len(data) > MAX_BYTES:
        raise ValueError("That file is over 5 MB. Upload a smaller PDF or Word file.")
    suffix = Path(filename).suffix.lower()
    is_pdf = suffix == ".pdf" and data.startswith(b"%PDF")
    if not is_pdf and not (suffix == ".docx" and data.startswith(b"PK")):
        raise ValueError("Upload a PDF or Word (.docx) file.")
    try:
        if is_pdf:
            # shortcut: a PDF whose streams unpack to gigabytes can still use a lot of memory; fine while the app runs
            # locally on the student's own file, cap stream sizes before deploying publicly.
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages[:MAX_PAGES])
        else:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                parts = z.infolist()
            xml = sum(p.file_size for p in parts if p.filename.endswith((".xml", ".rels")))
            if xml > MAX_XML or sum(p.file_size for p in parts) > MAX_UNPACKED:
                raise _TooLarge
            document = docx.Document(io.BytesIO(data))
            cells = (cell.text for table in document.tables for row in table.rows for cell in row.cells)
            text = "\n".join([*(p.text for p in document.paragraphs), *cells])
    except _TooLarge:
        raise ValueError("That file is too large to read. Upload a smaller PDF or Word file.") from None
    except Exception as error:  # broken files make the readers raise all sorts of errors, with internals in them
        raise ValueError("We couldn't open that file. Try saving it again as a PDF.") from error
    if not text.strip():
        raise ValueError("We couldn't find any text in that file. If it's a scanned image, upload a text-based PDF.")
    return text.strip()[:MAX_CHARS]
